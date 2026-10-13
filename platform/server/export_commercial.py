"""Render structured commercial content as inert text in DOCX/PDF documents."""
import json
import os
from pathlib import Path
import sys
from xml.sax.saxutils import escape

FONT_CANDIDATES = [
    ('/usr/share/fonts/truetype/nanum/NanumGothic.ttf', 'NanumGothic'),
    ('/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf', 'NanumBarunGothic'),
    ('/Library/Fonts/Arial Unicode.ttf', 'Arial Unicode MS'),
    ('/System/Library/Fonts/Supplemental/AppleGothic.ttf', 'AppleGothic'),
    ('/Library/Fonts/AppleGothic.ttf', 'AppleGothic'),
]


def clean(value):
    text = str(value or '')
    return ''.join(c for c in text if c in '\n\t\r' or (ord(c) >= 32 and ord(c) not in (0xfffe, 0xffff) and not 0xd800 <= ord(c) <= 0xdfff))


def metadata(data):
    template = data.get('template') or {}
    project = data.get('project') or {}
    status = '검토 승인본' if data.get('approved') is True else '검토 전 초안'
    if template.get('is_sample'):
        status += ' | 샘플 양식'
    return {
        'company': clean(template.get('company_name') or ''),
        'project': clean(project.get('name') or project.get('title') or ''),
        'revision': clean(data.get('revision_label') or ''),
        'footer': clean(template.get('footer') or ''),
        'status': status,
    }


def column_weights(table):
    weights = []
    for index, title in enumerate(table['columns']):
        lengths = [len(clean(row[index])) for row in table['rows']]
        lengths.sort()
        typical = lengths[int((len(lengths) - 1) * .8)] if lengths else 0
        weights.append(max(5, min(36, max(len(title), typical) ** .65 * 2.5)))
    total = sum(weights)
    return [weight / total for weight in weights]


def write_docx(data, output):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    landscape = any(len(table['columns']) > 6 for table in data['content']['tables'])
    section.page_width, section.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    font = next((name for filename, name in FONT_CANDIDATES if Path(filename).is_file()), 'NanumGothic')
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer'):
        style = doc.styles[name]
        style.font.name = font
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + attribute), None)
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        style.font.size = Pt(10.5 if name == 'Normal' else 9 if name in ('Header', 'Footer') else 23 if name == 'Title' else 14)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.35
    normal.space_after = Pt(7)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    meta = metadata(data)
    section.header.paragraphs[0].text = ' | '.join(filter(None, [meta['company'], meta['status']]))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(' | '.join(filter(None, [meta['revision'], meta['footer']])) + '  ')
    for field_name in ('PAGE', 'NUMPAGES'):
        if field_name == 'NUMPAGES': footer.add_run(' / ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), field_name)
        footer._p.append(field)
    doc.add_paragraph(clean(data['title']), 'Title')
    summary = ' | '.join(filter(None, [meta['project'], meta['revision'], meta['status']]))
    if summary: doc.add_paragraph(summary)
    for entry in data['content']['sections']:
        if entry['title']: doc.add_heading(clean(entry['title']), level=1)
        for paragraph in clean(entry['body']).split('\n'):
            doc.add_paragraph(paragraph)
    available = section.page_width - section.left_margin - section.right_margin
    for source in data['content']['tables']:
        if source['title']: doc.add_heading(clean(source['title']), level=1)
        table = doc.add_table(rows=1, cols=len(source['columns']))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [int(available * weight) for weight in column_weights(source)]
        for column, width in zip(table.columns, widths): column.width = width
        for index, title in enumerate(source['columns']): table.rows[0].cells[index].text = clean(title)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in source['rows']:
            cells = table.add_row().cells
            for index, value in enumerate(values): cells[index].text = clean(value)
        for row_index, row in enumerate(table.rows):
            # Keep ordinary records together; Word can still split a row taller than a page.
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell_index, cell in enumerate(row.cells):
                cell.width = widths[cell_index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    edge = OxmlElement('w:' + side)
                    for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: edge.set(qn('w:' + key), value)
                    borders.append(edge)
                properties.append(borders)
                padding = OxmlElement('w:tcMar')
                for side in ('top', 'left', 'bottom', 'right'):
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), '90')
                    margin.set(qn('w:type'), 'dxa')
                    padding.append(margin)
                properties.append(padding)
                if row_index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8EDF1')
                    properties.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.2
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.bold = row_index == 0
        doc.add_paragraph('')
    doc.core_properties.title = clean(data['title'])
    doc.core_properties.author = meta['company']
    doc.core_properties.subject = meta['status']
    doc.core_properties.comments = ''
    doc.save(output)


def write_pdf(data, output):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, CondPageBreak
    face = None
    for filename, _ in FONT_CANDIDATES:
        if Path(filename).is_file():
            try:
                pdfmetrics.registerFont(TTFont('Korean', filename))
                face = 'Korean'
                break
            except Exception:
                continue
    if face is None:
        face = 'HYSMyeongJo-Medium'
        pdfmetrics.registerFont(UnicodeCIDFont(face))
    page = landscape(A4) if any(len(t['columns']) > 6 for t in data['content']['tables']) else A4
    meta = metadata(data)
    doc = SimpleDocTemplate(output, pagesize=page, leftMargin=20*mm, rightMargin=20*mm,
        topMargin=26*mm, bottomMargin=24*mm, title=clean(data['title']), author=meta['company'])
    base = dict(fontName=face, fontSize=10, leading=15, textColor=colors.black, wordWrap='CJK', alignment=TA_LEFT)
    body = ParagraphStyle('Body', **base, spaceAfter=7)
    heading = ParagraphStyle('Heading', **{**base, 'fontSize': 14, 'leading': 20}, spaceBefore=14, spaceAfter=8, keepWithNext=True)
    table_heading = ParagraphStyle('TableHeading', parent=heading, keepWithNext=False)
    title = ParagraphStyle('Title', **{**base, 'fontSize': 23, 'leading': 31}, spaceAfter=13)
    cell = ParagraphStyle('Cell', **{**base, 'fontSize': 9, 'leading': 13}, spaceAfter=0)
    marginal = ParagraphStyle('Marginal', **{**base, 'fontSize': 8, 'leading': 11})
    def p(text, style=body): return Paragraph(escape(clean(text)).replace('\n', '<br/>') or '&#160;', style)
    story = [p(data['title'], title)]
    summary = ' | '.join(filter(None, [meta['project'], meta['revision'], meta['status']]))
    if summary: story.append(p(summary))
    for entry in data['content']['sections']:
        if entry['title']: story.append(p(entry['title'], heading))
        for paragraph in clean(entry['body']).split('\n'): story.append(p(paragraph))
    for source in data['content']['tables']:
        story.append(CondPageBreak(90))
        if source['title']: story.append(p(source['title'], table_heading))
        rows = [[p(value, cell) for value in source['columns']]]
        rows.extend([p(value, cell) for value in row] for row in source['rows'])
        table = LongTable(rows, colWidths=[doc.width * w for w in column_weights(source)], repeatRows=1, splitByRow=1, splitInRow=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8EDF1')),
            ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#D9D9D9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.extend([table, Spacer(1, 10)])
    def chrome(canvas, document):
        canvas.saveState()
        header = p(' | '.join(filter(None, [meta['company'], meta['status']])), marginal)
        _, height = header.wrap(doc.width, 40*mm)
        header.drawOn(canvas, doc.leftMargin, page[1] - 10*mm - height)
        footer = p(' | '.join(filter(None, [meta['revision'], meta['footer']])), marginal)
        _, height = footer.wrap(doc.width - 17*mm, 30*mm)
        footer.drawOn(canvas, doc.leftMargin, max(5*mm, 15*mm-height))
        canvas.setFont(face, 8)
        canvas.drawRightString(page[0]-doc.rightMargin, 11*mm, str(document.page))
        canvas.restoreState()
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)


def main():
    os.umask(0o077)
    if len(sys.argv) != 4 or sys.argv[3] not in ('docx', 'pdf'):
        raise ValueError('invalid arguments')
    source, output, fmt = sys.argv[1:]
    if Path(source).stat().st_size > 2*1024*1024:
        raise ValueError('input too large')
    data = json.loads(Path(source).read_text(encoding='utf-8'))
    (write_docx if fmt == 'docx' else write_pdf)(data, output)


if __name__ == '__main__':
    try:
        main()
    except ImportError:
        print(json.dumps({'ok': False, 'code': 'DEPENDENCY_MISSING'}), file=sys.stderr)
        sys.exit(1)
    except Exception:
        print(json.dumps({'ok': False, 'code': 'EXPORT_FAILED'}), file=sys.stderr)
        sys.exit(1)
